import os
import logging
import tempfile
import uuid
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from fastembed import TextEmbedding

from db import insert_document, insert_chunks, update_document_digest
from pdf.digest import generate_digest
from storage import StorageClient

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

_embedder: TextEmbedding | None = None


def get_embedder() -> TextEmbedding:
    global _embedder
    if _embedder is None:
        _embedder = TextEmbedding(
            model_name="sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
            max_length=512,
        )
    return _embedder


def extract_text_pdf(path: str) -> list[dict]:
    doc = fitz.open(path)
    pages = []
    for page_num, page in enumerate(doc, start=1):
        text = page.get_text()
        if text.strip():
            pages.append({"page_num": page_num, "text": text})
    doc.close()
    return pages


def _iter_docx_blocks(doc):
    """Yield paragraphs and tables in document order. python-docx exposes
    ``doc.paragraphs`` and ``doc.tables`` as separate collections, so reading
    only paragraphs silently drops every table."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _table_to_text(table) -> str:
    """Render a table as one row per line, cells joined by ' | '. The agenda
    and minuto-a-minuto schedules (hora/actividad/responsable) live entirely
    in tables; dropping them would strip every time, activity and speaker."""
    lines = []
    for row in table.rows:
        cells = [" ".join(c.text.split()) for c in row.cells]
        # Horizontally merged cells repeat the same text across grid columns.
        deduped: list[str] = []
        for c in cells:
            if c and (not deduped or deduped[-1] != c):
                deduped.append(c)
        if deduped:
            lines.append(" | ".join(deduped))
    return "\n".join(lines)


def extract_text_docx(path: str) -> list[dict]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(path)
    blocks: list[str] = []
    for item in _iter_docx_blocks(doc):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if text:
                blocks.append(text)
        elif isinstance(item, Table):
            text = _table_to_text(item)
            if text.strip():
                blocks.append(text)

    pages = []
    buffer = ""
    line = 1

    for block in blocks:
        if len(buffer) + len(block) < 3000:
            buffer += "\n\n" + block if buffer else block
        else:
            pages.append({"page_num": line, "text": buffer})
            line += 1
            buffer = block

    if buffer:
        pages.append({"page_num": line, "text": buffer})

    return pages


def chunk_pages(
    pages: list[dict], target_tokens: int = 800, overlap_words: int = 100
) -> list[dict]:
    chunks = []
    buffer = ""
    buffer_page = 1
    chunk_index = 0

    for page in pages:
        paragraphs = page["text"].split("\n\n")
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            words = para.split()
            para_len = len(words)

            if not buffer:
                buffer = para
                buffer_page = page["page_num"]
            elif len(buffer.split()) + para_len <= target_tokens:
                buffer += "\n\n" + para
            else:
                chunks.append(
                    {
                        "content": buffer,
                        "page_num": buffer_page,
                        "chunk_index": chunk_index,
                    }
                )
                chunk_index += 1
                overlap = buffer.split()[-overlap_words:] if len(buffer.split()) > overlap_words else []
                buffer = " ".join(overlap) + "\n\n" + para if overlap else para
                buffer_page = page["page_num"]

    if buffer.strip():
        chunks.append(
            {
                "content": buffer.strip(),
                "page_num": buffer_page,
                "chunk_index": chunk_index,
            }
        )

    return chunks


def embed_chunks(chunks: list[dict]) -> list[dict]:
    embedder = get_embedder()
    texts = [f"passage: {c['content']}" for c in chunks]
    embeddings = list(embedder.embed(texts))
    for chunk, emb in zip(chunks, embeddings):
        chunk["embedding"] = emb.tolist()
    return chunks


async def ingest_document(
    file_data: bytes,
    filename: str,
    storage: StorageClient,
    genai_client,
) -> dict:
    ext = Path(filename).suffix.lower()
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
    try:
        tmp.write(file_data)
        tmp_path = tmp.name
        tmp.close()

        if ext == ".pdf":
            pages = extract_text_pdf(tmp_path)
        else:
            pages = extract_text_docx(tmp_path)

        if not pages:
            return {"status": "error", "message": f"No text found in {ext}"}

        chunks = chunk_pages(pages)
        if not chunks:
            return {"status": "error", "message": "No chunks generated"}

        embed_chunks(chunks)

        file_uri = None
        try:
            gf = await genai_client.aio.files.upload(file=tmp_path)
            file_uri = gf.uri
        except Exception as e:
            logger.warning("[ingest] Gemini Files API upload failed (non-critical): %s", e)

        storage_path = f"{uuid.uuid4()}/{filename}"
        await storage.upload(storage_path, file_data)

        doc_id = await insert_document(
            name=filename,
            storage_path=storage_path,
            page_count=len(pages),
            file_uri=file_uri,
        )

        db_chunks = [
            {
                "document_id": doc_id,
                "content": c["content"],
                "page_num": c["page_num"],
                "chunk_index": c["chunk_index"],
                "embedding": c["embedding"],
            }
            for c in chunks
        ]
        await insert_chunks(db_chunks)

        # Digest failure is non-critical: retrieval still covers the document.
        full_text = "\n\n".join(p["text"] for p in pages)
        digest = await generate_digest(genai_client, filename, full_text)
        if digest:
            await update_document_digest(doc_id, digest)
            logger.info("[ingest] digest stored for '%s' (%d chars)", filename, len(digest))

        return {
            "status": "ready",
            "document_id": doc_id,
            "chunks_count": len(chunks),
            "pages": len(pages),
        }

    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
